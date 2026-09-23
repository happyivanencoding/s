# -*- coding: utf-8 -*-
"""
Modèle sectoriel Europe - version Python.

Objectif :
1. Lire uniquement les données mises à jour dans le fichier Excel FactSet.
2. Recalculer en Python les scores des piliers avec la même logique qu'Excel.
3. Produire les Top 3 / Worst 3 et la recommandation mensuelle.

Les commentaires sont volontairement simples et détaillés pour faciliter
les modifications futures.
"""

from pathlib import Path
import argparse
import math
import os

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string


# ---------------------------------------------------------------------------
# 1. PARAMÈTRES GÉNÉRAUX
# ---------------------------------------------------------------------------

def charger_fichier_excel_par_defaut():
    """
    Cherche le fichier Excel sans écrire de chemin local dans le code public.

    Ordre utilisé :
    1. variable d'environnement SCORE_SECTORIEL_EU_XLSM ;
    2. fichier local_config.py, présent uniquement sur la machine locale.
    """
    chemin_env = os.getenv("SCORE_SECTORIEL_EU_XLSM")
    if chemin_env:
        return chemin_env

    try:
        from local_config import FICHIER_EXCEL_EU
        return FICHIER_EXCEL_EU
    except ImportError:
        return None


FICHIER_EXCEL_PAR_DEFAUT = charger_fichier_excel_par_defaut()

DOSSIER_SORTIE_PAR_DEFAUT = Path(__file__).resolve().parent / "output"

# Ordre exact utilisé dans le fichier Excel Europe.
SECTEURS = [
    "Materials",
    "ConsStaples",
    "Pers. Goods",
    "Fin",
    "HealthCare",
    "Indus",
    "Oil",
    "Tech",
    "Telco",
    "Utili",
    "Travel & leisure",
    "Media",
    "Auto",
]

N_SECTEURS = len(SECTEURS)

# Excel utilise un historique de 60 mois + le mois courant = 61 observations.
FENETRE_HISTORIQUE = 60
N_OBSERVATIONS_RANK = FENETRE_HISTORIQUE + 1

# Nombre de secteurs retenus dans les Top / Worst.
N_TOP = 3
N_WORST = 3

# Deux sorties macro nécessaires dans la feuille "Cycle macro".
# M = régime macro : C / R / E / SD
# V = signal de taux : On / Off
COLONNE_CYCLE_MACRO = "M"
COLONNE_SIGNAL_TAUX = "V"

# Poids du modèle de base Europe : 6 piliers équipondérés.
POIDS_BASE = {
    "Leverage": 1 / 6,
    "Margin": 1 / 6,
    "Value": 1 / 6,
    "Momentum": 1 / 6,
    "Growth": 1 / 6,
    "Volatility": 1 / 6,
}

# Poids exacts de la table de tilt macro du fichier Europe.
POIDS_REGIME = {
    "C": {
        "Leverage": 0.10,
        "Margin": 0.00,
        "Value": 0.00,
        "Momentum": 0.30,
        "Growth": 0.20,
        "Volatility": 0.40,
    },
    "R": {
        "Leverage": 0.10,
        "Margin": 0.10,
        "Value": 0.10,
        "Momentum": 0.35,
        "Growth": 0.35,
        "Volatility": 0.00,
    },
    "E": {
        "Leverage": 1 / 6,
        "Margin": 1 / 6,
        "Value": 1 / 6,
        "Momentum": 1 / 6,
        "Growth": 1 / 6,
        "Volatility": 1 / 6,
    },
    "SD": {
        "Leverage": 0.00,
        "Margin": 0.00,
        "Value": 0.30,
        "Momentum": 0.10,
        "Growth": 0.30,
        "Volatility": 0.30,
    },
}

# Configuration des sous-variables directement utilisées dans les piliers.
# "ordre_rank":
#   1 = une valeur plus élevée est meilleure.
#   0 = une valeur plus faible est meilleure.
#
# "moyenne_sans_finance":
#   Excel exclut explicitement Financials de la moyenne sectorielle
#   pour certaines métriques non pertinentes.
VARIABLES_HISTORIQUES = {
    "Leverage": {
        "net_debt_ebitda": {
            "sheet": "Leverage_FMA",
            "colonne": "AF",
            "ordre_rank": 0,
            "moyenne_sans_finance": False,
            "mix_rank_transversal": False,
        },
        "fcf_total_debt": {
            "sheet": "Leverage_FMA",
            "colonne": "FB",
            "ordre_rank": 1,
            "moyenne_sans_finance": False,
            "mix_rank_transversal": False,
        },
        "debt_equity": {
            "sheet": "Leverage_FMA",
            "colonne": "GR",
            "ordre_rank": 0,
            "moyenne_sans_finance": False,
            "mix_rank_transversal": False,
        },
    },
    "Margin": {
        "operating_margin": {
            "sheet": "Margin_FMA",
            "colonne": "BV",
            "ordre_rank": 1,
            "moyenne_sans_finance": False,
            "mix_rank_transversal": True,
        },
        "net_margin": {
            "sheet": "Margin_FMA",
            "colonne": "DL",
            "ordre_rank": 1,
            "moyenne_sans_finance": False,
            "mix_rank_transversal": True,
        },
        "ebitda_margin": {
            "sheet": "Margin_FMA",
            "colonne": "GR",
            "ordre_rank": 1,
            "moyenne_sans_finance": False,
            "mix_rank_transversal": True,
        },
    },
    "Value": {
        "price_fcf": {
            "sheet": "Valuation_FMA_hist",
            "colonne": "DL",
            "ordre_rank": 0,
            "moyenne_sans_finance": True,
            "mix_rank_transversal": False,
            # Excel utilise 36 mois pour Technology sur cette métrique.
            "fenetre_par_secteur": {"Tech": 36},
        },
        "ev_ebitda": {
            "sheet": "Valuation_FMA_hist",
            "colonne": "FB",
            "ordre_rank": 0,
            "moyenne_sans_finance": True,
            "mix_rank_transversal": False,
            # Même exception de 36 mois pour Technology.
            "fenetre_par_secteur": {"Tech": 36},
        },
        "price_sales": {
            "sheet": "Valuation_FMA_hist",
            "colonne": "GR",
            "ordre_rank": 0,
            "moyenne_sans_finance": False,
            "mix_rank_transversal": False,
        },
    },
    "Growth": {
        "eps_growth": {
            "sheet": "Growth_FMA",
            "colonne": "AF",
            "ordre_rank": 1,
            "moyenne_sans_finance": False,
            "mix_rank_transversal": False,
        },
        "sales_growth": {
            "sheet": "Growth_FMA",
            "colonne": "DL",
            "ordre_rank": 1,
            "moyenne_sans_finance": False,
            "mix_rank_transversal": False,
        },
        "ebitda_growth": {
            "sheet": "Growth_FMA",
            "colonne": "FB",
            "ordre_rank": 1,
            "moyenne_sans_finance": True,
            "mix_rank_transversal": False,
        },
    },
}


# ---------------------------------------------------------------------------
# 2. PETITES FONCTIONS UTILITAIRES
# ---------------------------------------------------------------------------

def est_nombre(x):
    """Retourne True uniquement pour une valeur numérique exploitable."""
    return isinstance(x, (int, float, np.integer, np.floating)) and not pd.isna(x)


def convertir_date_excel(x):
    """Convertit une date Excel (numéro ou datetime) en Timestamp pandas."""
    if isinstance(x, (pd.Timestamp, np.datetime64)):
        return pd.Timestamp(x)

    # OpenPyXL renvoie souvent directement un datetime Python.
    if hasattr(x, "year") and hasattr(x, "month") and hasattr(x, "day"):
        return pd.Timestamp(x)

    if est_nombre(x):
        return pd.Timestamp("1899-12-30") + pd.to_timedelta(float(x), unit="D")

    return pd.NaT


def est_date_excel(x):
    """Teste si la cellule contient une date exploitable."""
    return not pd.isna(convertir_date_excel(x))


def rang_excel(valeur, valeurs, ordre):
    """
    Reproduit RANK d'Excel avec la méthode de rang "minimum".

    ordre = 1 : classement croissant, une grande valeur reçoit un grand rang.
    ordre = 0 : classement décroissant, une petite valeur reçoit un grand rang.
    """
    if pd.isna(valeur):
        return np.nan

    serie = pd.Series(valeurs, dtype="float64").dropna()
    if serie.empty:
        return np.nan

    # Excel compare les doubles directement. Une différence même minuscule
    # peut donc produire deux rangs différents.
    if ordre == 1:
        return 1 + int((serie < valeur).sum())
    return 1 + int((serie > valeur).sum())


def rang_transversal_0_10(ligne, ordre):
    """Transforme les 13 secteurs en rangs 0-10 comme dans Excel."""
    result = pd.Series(index=ligne.index, dtype=float)

    for secteur in ligne.index:
        valeur = ligne[secteur]
        if pd.isna(valeur):
            result[secteur] = np.nan
            continue
        rang = rang_excel(valeur, ligne.values, ordre)
        result[secteur] = 10 * rang / N_SECTEURS

    return result


def rang_transversal_momentum(ligne):
    """
    Rang 0-10 utilisé dans MOM_FMA.

    Excel écrit explicitement :
        RANK(...) / 13 * 10

    L'ordre des opérations est conservé car il peut laisser une différence
    de flottant de l'ordre de 1e-15 entre deux scores presque identiques.
    """
    result = pd.Series(index=ligne.index, dtype=float)

    for secteur in ligne.index:
        valeur = ligne[secteur]
        if pd.isna(valeur):
            result[secteur] = np.nan
            continue

        rang = rang_excel(valeur, ligne.values, ordre=1)
        result[secteur] = rang / N_SECTEURS * 10

    return result


def rang_secteurs(ligne):
    """
    Rang final des secteurs : 1 = Worst, 13 = Best.
    Le score d'entrée est déjà orienté "plus grand = meilleur".
    """
    result = pd.Series(index=ligne.index, dtype=float)

    for secteur in ligne.index:
        valeur = ligne[secteur]
        if pd.isna(valeur):
            result[secteur] = np.nan
            continue
        result[secteur] = rang_excel(valeur, ligne.values, ordre=1)

    return result


# ---------------------------------------------------------------------------
# 3. LECTURE DES DONNÉES EXCEL
# ---------------------------------------------------------------------------

def lire_dates_et_bloc(ws, colonne_depart, ligne_depart=8, colonne_date="A"):
    """
    Lit un bloc de 13 colonnes dans un FMA.
    Les lignes sont dans le même ordre que le fichier Excel :
    date la plus récente en premier.
    """
    col_start = column_index_from_string(colonne_depart)
    col_date = column_index_from_string(colonne_date)

    dates = []
    donnees = []
    ligne = ligne_depart

    while True:
        date_brute = ws.cell(ligne, col_date).value

        if not est_date_excel(date_brute):
            break

        valeurs = []
        for j in range(N_SECTEURS):
            v = ws.cell(ligne, col_start + j).value
            valeurs.append(float(v) if est_nombre(v) else np.nan)

        dates.append(convertir_date_excel(date_brute))
        donnees.append(valeurs)
        ligne += 1

    return pd.DataFrame(donnees, index=dates, columns=SECTEURS)


def lire_bloc_retours(ws):
    """
    Lit les rendements mensuels sectoriels dans Returns_EQ.
    Excel utilise P comme colonne de date et R:AD pour les 13 secteurs.
    """
    ligne = 7
    dates = []
    donnees = []

    while True:
        date_brute = ws.cell(ligne, column_index_from_string("P")).value
        if not est_date_excel(date_brute):
            break

        valeurs = []
        col_start = column_index_from_string("R")
        for j in range(N_SECTEURS):
            v = ws.cell(ligne, col_start + j).value
            valeurs.append(float(v) if est_nombre(v) else np.nan)

        dates.append(convertir_date_excel(date_brute))
        donnees.append(valeurs)
        ligne += 1

    return pd.DataFrame(donnees, index=dates, columns=SECTEURS)


def lire_cycle_macro(ws):
    """
    Lit directement les deux sorties macro déjà calculées dans Excel.

    Colonnes nécessaires dans la feuille "Cycle macro" :
    - M : régime macro = C / R / E / SD
    - V : signal de taux = On / Off

    Python ne recalcule ni le régime macro ni le signal de taux.
    """
    ligne = 4
    lignes = []

    col_cycle = column_index_from_string(COLONNE_CYCLE_MACRO)
    col_taux = column_index_from_string(COLONNE_SIGNAL_TAUX)

    while True:
        date_brute = ws.cell(ligne, 1).value
        if not est_date_excel(date_brute):
            break

        cycle = ws.cell(ligne, col_cycle).value
        signal_taux = ws.cell(ligne, col_taux).value

        lignes.append(
            {
                "date": convertir_date_excel(date_brute),
                "cycle": cycle if cycle in POIDS_REGIME else None,
                "signal_taux": signal_taux if signal_taux in {"On", "Off"} else None,
            }
        )
        ligne += 1

    return pd.DataFrame(lignes).set_index("date").sort_index()


# ---------------------------------------------------------------------------
# 4. CALCUL DES SOUS-VARIABLES HISTORIQUES
# ---------------------------------------------------------------------------

def calculer_diff_vs_moyenne(raw, moyenne_sans_finance=False):
    """
    Reproduit "diff vs. moy" :
    valeur du secteur - moyenne des secteurs du même mois.

    Pour certaines métriques Excel exclut Financials de la moyenne.
    """
    diff = pd.DataFrame(index=raw.index, columns=raw.columns, dtype=float)

    for date, ligne in raw.iterrows():
        base = ligne.copy()
        if moyenne_sans_finance:
            base = base.drop("Fin")

        moyenne = base.mean(skipna=True)
        diff.loc[date] = ligne - moyenne

    return diff


def calculer_score_historique(diff, ordre_rank, fenetre_par_secteur=None):
    """
    Reproduit la formule Excel :
    10 * RANK(valeur_courante, fenêtre_historique, ordre) / nombre_observations

    Par défaut Excel utilise 60 mois + le mois courant = 61 observations.
    Quelques cellules ont une fenêtre spécifique, par exemple Technology
    dans certaines métriques de Value.
    """
    fenetre_par_secteur = fenetre_par_secteur or {}
    score = pd.DataFrame(index=diff.index, columns=diff.columns, dtype=float)

    for i in range(len(diff)):
        for secteur in SECTEURS:
            valeur = diff.iloc[i][secteur]
            if pd.isna(valeur):
                continue

            n_mois = fenetre_par_secteur.get(secteur, FENETRE_HISTORIQUE)
            n_observations = n_mois + 1
            fin = i + n_observations

            if fin > len(diff):
                continue

            valeurs = diff.iloc[i:fin][secteur]
            rang = rang_excel(valeur, valeurs.values, ordre_rank)
            score.iloc[i, score.columns.get_loc(secteur)] = (
                10 * rang / n_observations
            )

    return score


def calculer_variable_historique(
    raw,
    ordre_rank,
    moyenne_sans_finance,
    mix_rank_transversal,
    fenetre_par_secteur=None,
):
    """Calcule une sous-variable exactement selon la logique FMA."""
    diff = calculer_diff_vs_moyenne(raw, moyenne_sans_finance)
    score_hist = calculer_score_historique(
        diff,
        ordre_rank,
        fenetre_par_secteur=fenetre_par_secteur,
    )

    if not mix_rank_transversal:
        return score_hist

    # Margin_FMA fait la moyenne entre :
    # - le score historique 5 ans ;
    # - le rang transversal du ratio brut au même mois.
    score = score_hist.copy()

    for date in score.index:
        score_cross = rang_transversal_0_10(raw.loc[date], ordre_rank)
        score.loc[date] = (score_hist.loc[date] + score_cross) / 2

    return score


# ---------------------------------------------------------------------------
# 5. PILIERS LEVERAGE / MARGIN / VALUE / GROWTH
# ---------------------------------------------------------------------------

def calculer_piliers_historiques(wb):
    """Calcule les 4 piliers basés sur les percentiles historiques."""
    sous_scores = {}
    piliers = {}

    for pilier, variables in VARIABLES_HISTORIQUES.items():
        scores_du_pilier = {}

        for nom_variable, cfg in variables.items():
            ws = wb[cfg["sheet"]]
            raw = lire_dates_et_bloc(ws, cfg["colonne"])

            score = calculer_variable_historique(
                raw=raw,
                ordre_rank=cfg["ordre_rank"],
                moyenne_sans_finance=cfg["moyenne_sans_finance"],
                mix_rank_transversal=cfg["mix_rank_transversal"],
                fenetre_par_secteur=cfg.get("fenetre_par_secteur"),
            )

            sous_scores[nom_variable] = score
            scores_du_pilier[nom_variable] = score

        # Même logique que AVERAGE dans les formules principales.
        # Pour Finance, Growth utilise seulement EPS + Sales car EBITDA est N/A.
        index_commun = next(iter(scores_du_pilier.values())).index
        pilier_df = pd.DataFrame(index=index_commun, columns=SECTEURS, dtype=float)

        for date in index_commun:
            for secteur in SECTEURS:
                valeurs = []

                for nom_variable, score in scores_du_pilier.items():
                    if pilier == "Growth" and secteur == "Fin" and nom_variable == "ebitda_growth":
                        continue
                    if pilier == "Value" and secteur == "Fin" and nom_variable == "price_fcf":
                        continue

                    # Dans Excel, Technology utilise uniquement P/FCF + EV/EBITDA.
                    # Price/Sales est explicitement exclu de la formule X8.
                    if pilier == "Value" and secteur == "Tech" and nom_variable == "price_sales":
                        continue

                    v = score.at[date, secteur]
                    valeurs.append(v)

                # Excel propage une erreur si une composante explicitement utilisée est N/A.
                if any(pd.isna(v) for v in valeurs):
                    pilier_df.at[date, secteur] = np.nan
                else:
                    # AVERAGE Excel : addition simple puis division.
                    # On évite np.mean pour garder les mêmes arrondis binaires.
                    pilier_df.at[date, secteur] = sum(valeurs) / len(valeurs)

        # Growth contient dans Excel quelques ex-aequo exacts qui peuvent
        # devenir différents de ~1e-15 après les additions Python.
        # Un arrondi à 14 décimales restitue le même classement Excel,
        # sans modifier les autres piliers où ces micro-écarts sont utilisés.
        if pilier == "Growth":
            pilier_df = pilier_df.round(14)

        piliers[pilier] = pilier_df

    return piliers, sous_scores


# ---------------------------------------------------------------------------
# 6. PILIER MOMENTUM
# ---------------------------------------------------------------------------

def calculer_momentum(wb):
    """Reproduit MOM_FMA avec ses trois composantes."""
    ws = wb["MOM_FMA"]

    prix = lire_dates_et_bloc(ws, "AT")
    up = lire_dates_et_bloc(ws, "DL")
    down = lire_dates_et_bloc(ws, "DZ")
    unchanged = lire_dates_et_bloc(ws, "EN")

    ret_6m_1m = pd.DataFrame(index=prix.index, columns=SECTEURS, dtype=float)
    ret_12m_1m = pd.DataFrame(index=prix.index, columns=SECTEURS, dtype=float)

    for i in range(len(prix)):
        # Excel : AT9 / AT14 - 1
        if i + 6 < len(prix):
            ret_6m_1m.iloc[i] = prix.iloc[i + 1] / prix.iloc[i + 6] - 1

        # Excel : AT9 / AT20 - 1
        if i + 12 < len(prix):
            ret_12m_1m.iloc[i] = prix.iloc[i + 1] / prix.iloc[i + 12] - 1

        # Exception présente dans le fichier Excel :
        # Technology utilise le mois courant comme numérateur
        # (BA8/BA14 et BA8/BA20), alors que les autres secteurs
        # excluent le dernier mois.
        if i + 6 < len(prix):
            ret_6m_1m.at[prix.index[i], "Tech"] = (
                prix.iloc[i]["Tech"] / prix.iloc[i + 6]["Tech"] - 1
            )
        if i + 12 < len(prix):
            ret_12m_1m.at[prix.index[i], "Tech"] = (
                prix.iloc[i]["Tech"] / prix.iloc[i + 12]["Tech"] - 1
            )

    revision_ratio = (up - down) / (up + down + unchanged)

    score_6m = ret_6m_1m.apply(rang_transversal_momentum, axis=1)
    score_12m = ret_12m_1m.apply(rang_transversal_momentum, axis=1)
    score_revision = revision_ratio.apply(rang_transversal_momentum, axis=1)

    # On reproduit AVERAGE cellule par cellule.
    # Cela conserve les mêmes micro-différences de flottants qu'Excel,
    # qui peuvent parfois influencer RANK sur deux scores quasi identiques.
    momentum = pd.DataFrame(
        index=score_6m.index,
        columns=SECTEURS,
        dtype=float,
    )

    for date in momentum.index:
        for secteur in SECTEURS:
            valeurs = [
                score_6m.at[date, secteur],
                score_12m.at[date, secteur],
                score_revision.at[date, secteur],
            ]

            if any(pd.isna(v) for v in valeurs):
                momentum.at[date, secteur] = np.nan
            else:
                momentum.at[date, secteur] = sum(valeurs) / 3

    sous_scores = {
        "momentum_6m_1m": score_6m,
        "momentum_12m_1m": score_12m,
        "earnings_revision_ratio": score_revision,
    }

    return momentum, sous_scores


# ---------------------------------------------------------------------------
# 7. PILIER LOW VOLATILITY
# ---------------------------------------------------------------------------

def calculer_volatilite(wb):
    """Reproduit Vol_FMA à partir des rendements mensuels de Returns_EQ."""
    retours = lire_bloc_retours(wb["Returns_EQ"])

    vol_6m = pd.DataFrame(index=retours.index, columns=SECTEURS, dtype=float)
    vol_down_18m = pd.DataFrame(index=retours.index, columns=SECTEURS, dtype=float)

    for i in range(len(retours)):
        # Excel : STDEVA(current:OFFSET(current,6))*SQRT(12)
        # OFFSET(...,6) inclut 7 observations : courant + 6 lignes.
        if i + 7 <= len(retours):
            fenetre = retours.iloc[i : i + 7]
            vol_6m.iloc[i] = fenetre.std(ddof=1) * math.sqrt(12)

        # Excel : STDEVA(IF(return<0, return))*SQRT(12)
        # Dans cette formule matricielle, les FALSE sont ignorés :
        # on calcule donc l'écart-type uniquement sur les mois négatifs.
        if i + 19 <= len(retours):
            fenetre = retours.iloc[i : i + 19]

            for secteur in SECTEURS:
                negatifs = fenetre[secteur][fenetre[secteur] < 0].dropna()
                if len(negatifs) >= 2:
                    vol_down_18m.at[retours.index[i], secteur] = (
                        negatifs.std(ddof=1) * math.sqrt(12)
                    )

    # Les deux sous-variables utilisent la même logique historique :
    # faible volatilité = meilleur signal => ordre RANK = 0.
    score_vol = calculer_variable_historique(
        vol_6m,
        ordre_rank=0,
        moyenne_sans_finance=False,
        mix_rank_transversal=False,
    )
    score_down = calculer_variable_historique(
        vol_down_18m,
        ordre_rank=0,
        moyenne_sans_finance=False,
        mix_rank_transversal=False,
    )

    volatility = (score_vol + score_down) / 2

    sous_scores = {
        "volatility_6m": score_vol,
        "downside_volatility_18m": score_down,
    }

    return volatility, sous_scores, retours


# ---------------------------------------------------------------------------
# 8. AGRÉGATION DES PILIERS ET RECOMMANDATIONS
# ---------------------------------------------------------------------------

def aligner_piliers(piliers):
    """Aligne les six piliers sur les dates communes."""
    dates = None
    for df in piliers.values():
        dates = df.index if dates is None else dates.intersection(df.index)

    dates = dates.sort_values(ascending=False)
    return {nom: df.loc[dates].copy() for nom, df in piliers.items()}


def calculer_rangs_piliers(piliers):
    """Transforme chaque score 0-10 en rang sectoriel 1-13."""
    rangs = {}

    for nom, df in piliers.items():
        rang_df = pd.DataFrame(index=df.index, columns=SECTEURS, dtype=float)
        for date in df.index:
            rang_df.loc[date] = rang_secteurs(df.loc[date])
        rangs[nom] = rang_df

    return rangs


def calculer_score_pondere(rangs, poids, date):
    """Somme pondérée des rangs des piliers pour une date."""
    resultat = pd.Series(0.0, index=SECTEURS)
    valide = pd.Series(True, index=SECTEURS)

    for pilier, poids_pilier in poids.items():
        serie = rangs[pilier].loc[date]
        valide &= serie.notna()
        resultat += serie.fillna(0) * poids_pilier

    resultat[~valide] = np.nan
    return resultat


def choisir_top_worst(top_count, bottom_count, rang_global):
    """
    Remplacement transparent de la fonction XLL tab_invest_eu.

    Logique visible dans Excel :
    - un secteur Bottom >= 3 ne peut pas être Positive ;
    - un secteur Top >= 3 ne peut pas être Negative ;
    - priorité au nombre de confirmations ;
    - le rang global sert de départage.
    """
    tableau = pd.DataFrame(
        {
            "top_count": top_count,
            "bottom_count": bottom_count,
            "rang_global": rang_global,
        }
    )

    candidats_top = tableau[tableau["bottom_count"] < N_WORST].copy()
    candidats_top = candidats_top.sort_values(
        ["top_count", "rang_global", "bottom_count"],
        ascending=[False, False, True],
    )
    top = list(candidats_top.head(N_TOP).index)

    candidats_worst = tableau[tableau["top_count"] < N_TOP].copy()
    candidats_worst = candidats_worst.sort_values(
        ["bottom_count", "rang_global", "top_count"],
        ascending=[False, True, True],
    )
    worst = list(candidats_worst.head(N_WORST).index)

    return top, worst


def calculer_resultats_finaux(piliers, cycle_macro):
    """Calcule l'historique complet du modèle final Europe."""
    piliers = aligner_piliers(piliers)
    rangs = calculer_rangs_piliers(piliers)

    lignes = []

    for date in piliers["Leverage"].index:
        # Le modèle final a besoin du régime macro.
        if date not in cycle_macro.index:
            continue

        regime = cycle_macro.at[date, "cycle"]
        signal_taux = cycle_macro.at[date, "signal_taux"]

        if regime not in POIDS_REGIME:
            continue
        if signal_taux not in {"On", "Off"}:
            continue

        # Score 5F de base : six piliers équipondérés.
        score_base = calculer_score_pondere(rangs, POIDS_BASE, date)
        rang_base = rang_secteurs(score_base)

        # Score Tilt : poids dépendants du régime macro.
        score_tilt = calculer_score_pondere(rangs, POIDS_REGIME[regime], date)
        rang_tilt = rang_secteurs(score_tilt)

        # "+ macro" dans Excel :
        # le 7e pilier est le rang issu du modèle Tilt.
        score_macro = pd.Series(0.0, index=SECTEURS)
        valide_macro = pd.Series(True, index=SECTEURS)

        for pilier in POIDS_BASE:
            serie = rangs[pilier].loc[date]
            valide_macro &= serie.notna()
            score_macro += serie.fillna(0) / 7

        valide_macro &= rang_tilt.notna()
        score_macro += rang_tilt.fillna(0) / 7
        score_macro[~valide_macro] = np.nan

        rang_global = rang_secteurs(score_macro)

        # Comptage Top 3 / Worst 3.
        top_count = pd.Series(0, index=SECTEURS, dtype=int)
        bottom_count = pd.Series(0, index=SECTEURS, dtype=int)

        for pilier in ["Leverage", "Margin", "Value", "Momentum", "Growth", "Volatility"]:
            r = rangs[pilier].loc[date]
            top_count += (r > N_SECTEURS - N_TOP).fillna(False).astype(int)
            bottom_count += (r <= N_WORST).fillna(False).astype(int)

        # Le "macro pillar" compte deux fois dans Excel.
        top_count += 2 * (rang_tilt > N_SECTEURS - N_TOP).fillna(False).astype(int)
        bottom_count += 2 * (rang_tilt <= N_WORST).fillna(False).astype(int)

        # Si le signal de taux est On, Value et Leverage comptent une fois de plus.
        if signal_taux == "On":
            for pilier in ["Leverage", "Value"]:
                r = rangs[pilier].loc[date]
                top_count += (r > N_SECTEURS - N_TOP).fillna(False).astype(int)
                bottom_count += (r <= N_WORST).fillna(False).astype(int)

        # Exception explicite du fichier Excel :
        # les formules Top et Worst de Travel & leisure sont multipliées par 0.
        top_count["Travel & leisure"] = 0
        bottom_count["Travel & leisure"] = 0

        top, worst = choisir_top_worst(top_count, bottom_count, rang_global)

        for secteur in SECTEURS:
            if secteur in top:
                reco = "Positive"
            elif secteur in worst:
                reco = "Negative"
            else:
                reco = "Neutral"

            ligne = {
                "date": date,
                "secteur": secteur,
                "cycle": regime,
                "signal_taux": signal_taux,
                "score_base": score_base[secteur],
                "rang_base": rang_base[secteur],
                "score_tilt": score_tilt[secteur],
                "rang_tilt_macro": rang_tilt[secteur],
                "score_global_macro": score_macro[secteur],
                "rang_global": rang_global[secteur],
                "top_count": int(top_count[secteur]),
                "bottom_count": int(bottom_count[secteur]),
                "recommendation": reco,
            }

            for pilier in rangs:
                ligne[f"score_{pilier.lower()}"] = piliers[pilier].at[date, secteur]
                ligne[f"rang_{pilier.lower()}"] = rangs[pilier].at[date, secteur]

            lignes.append(ligne)

    return pd.DataFrame(lignes), piliers, rangs, cycle_macro


# ---------------------------------------------------------------------------
# 9. FONCTION PRINCIPALE
# ---------------------------------------------------------------------------

def calculer_modele(
    fichier_excel=None,
):
    """Point d'entrée principal réutilisé par les scripts de plot et backtest."""
    if fichier_excel is None:
        fichier_excel = FICHIER_EXCEL_PAR_DEFAUT

    if fichier_excel is None:
        raise ValueError(
            "Aucun fichier Excel configuré. "
            "Utiliser --excel, SCORE_SECTORIEL_EU_XLSM ou local_config.py."
        )

    fichier_excel = Path(fichier_excel)

    if not fichier_excel.exists():
        raise FileNotFoundError(f"Fichier introuvable : {fichier_excel}")

    # data_only=True : on lit les valeurs mises en cache après le refresh FactSet.
    wb = load_workbook(
        fichier_excel,
        data_only=True,
        read_only=False,
        keep_vba=False,
    )

    try:
        piliers, sous_scores = calculer_piliers_historiques(wb)

        momentum, sous_momentum = calculer_momentum(wb)
        volatility, sous_vol, retours = calculer_volatilite(wb)

        piliers["Momentum"] = momentum
        piliers["Volatility"] = volatility

        sous_scores.update(sous_momentum)
        sous_scores.update(sous_vol)

        cycle_macro = lire_cycle_macro(wb["Cycle macro"])

        historique, piliers, rangs, cycle_macro = calculer_resultats_finaux(
            piliers,
            cycle_macro,
        )

        return {
            "historique": historique,
            "piliers": piliers,
            "rangs": rangs,
            "sous_scores": sous_scores,
            "retours": retours,
            "cycle_macro": cycle_macro,
        }
    finally:
        wb.close()


def sauvegarder_sorties(resultats, dossier_sortie):
    """Sauvegarde des CSV simples et faciles à contrôler."""
    dossier = Path(dossier_sortie)
    dossier.mkdir(parents=True, exist_ok=True)

    historique = resultats["historique"].copy()
    historique.to_csv(dossier / "eu_historique_modele.csv", index=False)

    # Les six piliers peuvent être plus récents que le régime macro.
    # On sauvegarde donc toujours leur dernière date disponible.
    piliers = resultats["piliers"]
    rangs = resultats["rangs"]

    dates_piliers = None
    for df in piliers.values():
        dates_piliers = (
            df.index
            if dates_piliers is None
            else dates_piliers.intersection(df.index)
        )

    if dates_piliers is not None and len(dates_piliers):
        date_piliers = dates_piliers.max()
        lignes_piliers = []

        for secteur in SECTEURS:
            ligne = {
                "date": date_piliers,
                "secteur": secteur,
            }

            for pilier in [
                "Leverage",
                "Margin",
                "Value",
                "Momentum",
                "Growth",
                "Volatility",
            ]:
                ligne[f"score_{pilier.lower()}"] = piliers[pilier].at[
                    date_piliers, secteur
                ]
                ligne[f"rang_{pilier.lower()}"] = rangs[pilier].at[
                    date_piliers, secteur
                ]

            lignes_piliers.append(ligne)

        pd.DataFrame(lignes_piliers).to_csv(
            dossier / "eu_piliers_latest_available.csv",
            index=False,
        )

    if historique.empty:
        return

    date_latest = historique["date"].max()
    latest = historique[historique["date"] == date_latest].copy()
    latest = latest.sort_values("rang_global", ascending=False)
    latest.to_csv(dossier / "eu_recommandations_latest.csv", index=False)

    colonnes = [
        "date",
        "secteur",
        "score_leverage",
        "score_margin",
        "score_value",
        "score_momentum",
        "score_growth",
        "score_volatility",
        "rang_global",
        "recommendation",
    ]
    latest[colonnes].to_csv(dossier / "eu_piliers_latest_with_reco.csv", index=False)


def afficher_latest(resultats):
    """Affichage console volontairement compact."""
    historique = resultats["historique"]

    if historique.empty:
        print("Aucun résultat calculé.")
        return

    date_latest = historique["date"].max()
    latest = historique[historique["date"] == date_latest].copy()
    latest = latest.sort_values("rang_global", ascending=False)

    cols = [
        "secteur",
        "rang_global",
        "top_count",
        "bottom_count",
        "recommendation",
        "cycle",
        "signal_taux",
    ]

    print(f"\nDate de la recommandation finale : {date_latest.date()}")
    print(latest[cols].to_string(index=False))

    # Information importante si le fichier macro est moins récent.
    dates_piliers = None
    for df in resultats["piliers"].values():
        dates_piliers = (
            df.index
            if dates_piliers is None
            else dates_piliers.intersection(df.index)
        )

    if dates_piliers is not None and len(dates_piliers):
        date_piliers = dates_piliers.max()
        if date_piliers > date_latest:
            print(
                f"\nAttention : les piliers sont disponibles jusqu'au "
                f"{date_piliers.date()}, mais le régime macro permet une "
                f"recommandation finale seulement jusqu'au {date_latest.date()}."
            )


def main():
    parser = argparse.ArgumentParser(description="Modèle sectoriel Europe")
    parser.add_argument(
        "--excel",
        default=FICHIER_EXCEL_PAR_DEFAUT,
        help="Chemin du fichier Score_Sectoriel_EU.xlsm",
    )
    parser.add_argument(
        "--output",
        default=str(DOSSIER_SORTIE_PAR_DEFAUT),
        help="Dossier de sortie",
    )
    args = parser.parse_args()

    resultats = calculer_modele(args.excel)
    sauvegarder_sorties(resultats, args.output)
    afficher_latest(resultats)


if __name__ == "__main__":
    main()
