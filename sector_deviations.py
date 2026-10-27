"""Copie les recommandations dans le classeur d'allocation existant.

Le classeur conserve son mapping, ses formules et sa présentation.
"""

from pathlib import Path
from tempfile import TemporaryDirectory
import os
import shutil

import pandas as pd
from openpyxl.utils import get_column_letter

from local_config import EU_SECTORS, US_SECTORS, SECTOR_DEVIATION_EXCEL_FILE


SECTORS = {"EU": EU_SECTORS, "US": US_SECTORS}


def allocation_signals(history, market):
    """Prépare une ligne de recommandations par date de signal."""
    sectors = SECTORS[market]
    table = history.pivot(index="date", columns="secteur", values="recommendation").sort_index()
    if not set(sectors) <= set(table.columns) or table[sectors].isna().any().any():
        raise ValueError(f"Recommandations incomplètes pour {market}.")
    return table[sectors]


def excel_date(value):
    """Convertit une date en numéro de série Excel pour Value2."""
    return (pd.Timestamp(value) - pd.Timestamp("1899-12-30")).total_seconds() / 86400


def _write_signals(workbook, history, market):
    signals = allocation_signals(history, market)
    sectors = SECTORS[market]
    sheet = workbook.Worksheets(f"Signals_{market}")
    allocation = workbook.Worksheets(f"secto_{market.lower()}")
    signal_column = get_column_letter(len(sectors) + 2)
    formula_column = get_column_letter(len(sectors) + 6)

    headers = sheet.Range(f"C1:{signal_column}1").Value2[0]
    if tuple(headers) != tuple(sectors):
        raise ValueError(f"Colonnes de signaux incompatibles pour {market}.")

    last_signal_row = sheet.Cells(sheet.Rows.Count, 1).End(-4162).Row
    last_allocation_row = allocation.Cells(allocation.Rows.Count, 1).End(-4162).Row
    if last_signal_row < 2 or last_allocation_row < 2:
        raise ValueError(f"Le classeur ne contient pas les formules existantes pour {market}.")

    rows = [tuple(row) for row in sheet.Range(f"A2:{signal_column}{last_signal_row}").Value2]
    dates = [row[0] for row in rows]
    if any(not isinstance(value, (int, float)) for value in dates) or dates != sorted(set(dates)):
        raise ValueError(f"Dates de signaux invalides ou non classées pour {market}.")
    positions = {date: index for index, date in enumerate(dates)}
    codes = {"Positive": 1, "Negative": -1, "Neutral": 0}

    for date, values in signals.iterrows():
        signal_date = excel_date(date)
        allocation_date = excel_date(date + pd.offsets.MonthBegin(1))
        entry = (signal_date, allocation_date, *[codes[value] for value in values])
        if signal_date in positions:
            rows[positions[signal_date]] = entry
        elif rows and signal_date < rows[-1][0]:
            raise ValueError(f"Mois historique absent du classeur {market} : {date:%d/%m/%Y}.")
        else:
            positions[signal_date] = len(rows)
            rows.append(entry)

    last_row = len(rows) + 1
    # Excel prolonge les formules et la présentation de la dernière ligne.
    if last_row > last_signal_row:
        sheet.Range(f"A{last_signal_row}:{formula_column}{last_row}").FillDown()
    if last_row > last_allocation_row:
        allocation.Range(f"A{last_allocation_row}:U{last_row}").FillDown()
    sheet.Range(f"A2:{signal_column}{last_row}").Value2 = tuple(rows)
    return last_row


def update_sector_deviations(models, output_path=SECTOR_DEVIATION_EXCEL_FILE):
    """Actualise les signaux EU/US sans reconstruire le classeur."""
    import win32com.client

    by_market = {model.universe: model for model in models}
    if set(by_market) != {"EU", "US"}:
        raise ValueError("Les deux marchés EU et US sont nécessaires.")
    target = Path(output_path).expanduser().resolve()
    if not target.is_file():
        raise FileNotFoundError(f"Classeur d'allocation introuvable : {target}")
    if (target.parent / ("~$" + target.name)).exists():
        raise RuntimeError(f"Fermer le fichier d'allocation : {target.name}")

    with TemporaryDirectory(prefix="sector_deviations_", dir=target.parent) as temporary:
        working = Path(temporary) / target.name
        shutil.copy2(target, working)
        app = win32com.client.DispatchEx("Excel.Application")
        app.Visible = False
        app.DisplayAlerts = False
        app.EnableEvents = False
        app.AutomationSecurity = 3
        app.AskToUpdateLinks = False
        workbook = None
        try:
            workbook = app.Workbooks.Open(str(working), UpdateLinks=0, ReadOnly=False)
            if workbook.ReadOnly:
                raise RuntimeError("Classeur d'allocation ouvert en lecture seule.")

            last_rows = {}
            for market in ("US", "EU"):
                last_rows[market] = _write_signals(workbook, by_market[market].results["history"], market)
            app.CalculateFull()
            for market, last_row in last_rows.items():
                sheet = workbook.Worksheets(f"secto_{market.lower()}")
                totals = sheet.Range(f"U2:U{last_row}").Value2
                if not isinstance(totals, tuple):
                    totals = ((totals,),)
                if any(not isinstance(row[0], (int, float)) or abs(row[0]) > 1e-10 for row in totals):
                    raise ValueError(f"Déviations non nulles ou indisponibles dans secto_{market.lower()}.")
            workbook.Save()
        finally:
            if workbook is not None:
                workbook.Close(False)
            app.Quit()

        os.replace(working, target)
    return target
