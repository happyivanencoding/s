# -*- coding: utf-8 -*-
"""Entrées data des modèles sectoriels."""

from pathlib import Path
import os

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from local_config import (
    CONFIG_HISTORIQUE,
    FICHIER_EXCEL_MACRO,
    MODELE_EU,
    MODELE_US,
)

try:
    from local_config_private import (
        FICHIER_EXCEL_EU as FICHIER_EXCEL_EU_PRIVE,
        FICHIER_EXCEL_US as FICHIER_EXCEL_US_PRIVE,
        FICHIER_EXCEL_MACRO as FICHIER_EXCEL_MACRO_PRIVE,
    )
except ImportError:
    FICHIER_EXCEL_EU_PRIVE = ""
    FICHIER_EXCEL_US_PRIVE = ""
    FICHIER_EXCEL_MACRO_PRIVE = ""


def choisir_chemin(configuration, surcharge, variable_env):
    """Sélectionne le chemin configuré selon l'ordre de priorité défini."""
    chemin_env = os.getenv(variable_env)

    if chemin_env:
        return chemin_env
    if surcharge:
        return surcharge
    if configuration:
        return configuration

    return None


FICHIER_EXCEL_EU_PAR_DEFAUT = choisir_chemin(
    MODELE_EU["fichier_excel"],
    FICHIER_EXCEL_EU_PRIVE,
    MODELE_EU["env_excel"],
)

FICHIER_EXCEL_US_PAR_DEFAUT = choisir_chemin(
    MODELE_US["fichier_excel"],
    FICHIER_EXCEL_US_PRIVE,
    MODELE_US["env_excel"],
)

FICHIER_EXCEL_PAR_DEFAUT = FICHIER_EXCEL_EU_PAR_DEFAUT

FICHIER_MACRO_PAR_DEFAUT = choisir_chemin(
    FICHIER_EXCEL_MACRO,
    FICHIER_EXCEL_MACRO_PRIVE,
    "SCORE_MACRO_EU_XLSX",
)


def ouvrir_workbooks(
    fichier_excel=None,
    fichier_macro=None,
    fichier_excel_defaut=None,
):
    """Ouvre le fichier sectoriel et le fichier macro."""
    fichier_excel = (
        fichier_excel
        or fichier_excel_defaut
        or FICHIER_EXCEL_EU_PAR_DEFAUT
    )
    fichier_macro = fichier_macro or FICHIER_MACRO_PAR_DEFAUT

    if not fichier_excel:
        raise ValueError("Aucun fichier sectoriel configuré.")
    if not fichier_macro:
        raise ValueError("Aucun fichier macro configuré.")

    fichier_excel = Path(fichier_excel)
    fichier_macro = Path(fichier_macro)

    if not fichier_excel.exists():
        raise FileNotFoundError(
            f"Fichier sectoriel introuvable : {fichier_excel}"
        )
    if not fichier_macro.exists():
        raise FileNotFoundError(
            f"Fichier macro introuvable : {fichier_macro}"
        )

    wb_secteur = load_workbook(
        fichier_excel,
        data_only=True,
        read_only=False,
        keep_vba=False,
    )
    wb_macro = load_workbook(
        fichier_macro,
        data_only=True,
        read_only=False,
        keep_vba=False,
    )

    return wb_secteur, wb_macro


def est_nombre(x):
    """Indique si une cellule contient une valeur numérique exploitable."""
    return isinstance(
        x,
        (int, float, np.integer, np.floating),
    ) and not pd.isna(x)


def convertir_date_excel(x):
    """Convertit une date Excel ou Python en Timestamp pandas."""
    if isinstance(x, (pd.Timestamp, np.datetime64)):
        return pd.Timestamp(x)

    if hasattr(x, "year") and hasattr(x, "month") and hasattr(x, "day"):
        return pd.Timestamp(x)

    if est_nombre(x):
        return pd.Timestamp("1899-12-30") + pd.to_timedelta(
            float(x),
            unit="D",
        )

    return pd.NaT


def est_date_excel(x):
    """Vérifie qu'une cellule contient une date exploitable."""
    date = convertir_date_excel(x)

    if pd.isna(date):
        return False

    return 1980 <= date.year <= 2100


def normaliser_date_mensuelle(x):
    """Utilise le dernier jour calendaire du mois comme clé de date."""
    date = convertir_date_excel(x)

    if pd.isna(date):
        return pd.NaT

    return pd.Timestamp(date).to_period("M").to_timestamp("M")


def normaliser_index_mensuel(df):
    """Normalise l'index en fins de mois et conserve une ligne par mois."""
    resultat = df.copy()
    resultat.index = pd.DatetimeIndex(
        [normaliser_date_mensuelle(x) for x in resultat.index]
    )
    resultat = resultat[
        ~resultat.index.duplicated(keep="first")
    ]
    return resultat.sort_index(ascending=False)


def chemin_historique():
    """Retourne le fichier Parquet qui contient tout l'historique."""
    return (
        Path(__file__).resolve().parent
        / CONFIG_HISTORIQUE["fichier"]
    )


def figer_historique(df, dates_source, cle):
    """
    Conserve la première observation enregistrée pour chaque mois.
    Une révision ultérieure du fichier source ne remplace pas le mois existant.
    """
    resultat = normaliser_index_mensuel(df)

    dates_source = pd.Series(
        [convertir_date_excel(x) for x in dates_source],
        index=[
            normaliser_date_mensuelle(x)
            for x in dates_source
        ],
        name="date_source",
    )
    dates_source = dates_source[
        ~dates_source.index.duplicated(keep="first")
    ]

    courant = resultat.copy()
    courant.insert(
        0,
        "date_source",
        dates_source.reindex(courant.index).values,
    )
    courant.insert(0, "date", courant.index)
    courant.insert(0, "series_key", cle)
    courant = courant.reset_index(drop=True)

    colonnes_data = list(df.columns)
    courant = courant[
        courant[colonnes_data]
        .notna()
        .any(axis=1)
    ]

    if not CONFIG_HISTORIQUE.get("actif", True):
        return resultat

    fichier = chemin_historique()

    if fichier.exists():
        historique = pd.read_parquet(fichier)
        historique["date"] = pd.to_datetime(
            historique["date"]
        )
        historique["date_source"] = pd.to_datetime(
            historique["date_source"]
        )

        dates_existantes = historique.loc[
            historique["series_key"] == cle,
            "date",
        ]

        nouveaux_mois = courant.loc[
            ~courant["date"].isin(dates_existantes)
        ]

        if len(nouveaux_mois):
            historique = pd.concat(
                [historique, nouveaux_mois],
                ignore_index=True,
                sort=False,
            )
            historique = historique.sort_values(
                ["series_key", "date"],
                ascending=[True, False],
            )
            historique.to_parquet(
                fichier,
                index=False,
            )
    else:
        historique = courant.sort_values(
            ["series_key", "date"],
            ascending=[True, False],
        )
        historique.to_parquet(
            fichier,
            index=False,
        )

    serie = historique[
        historique["series_key"] == cle
    ].copy()

    serie = serie.sort_values(
        "date",
        ascending=False,
    ).set_index("date")

    colonnes = list(df.columns)

    for colonne in colonnes:
        if colonne not in serie.columns:
            serie[colonne] = np.nan

    return serie[colonnes].copy()


def lire_dates_et_bloc(
    ws,
    colonne_depart,
    secteurs=None,
    cle_prefixe="",
    ligne_depart=8,
    colonne_date="A",
):
    """Lit un bloc sectoriel puis fige son historique mensuel."""
    secteurs = secteurs or MODELE_EU["secteurs"]

    col_start = column_index_from_string(colonne_depart)
    col_date = column_index_from_string(colonne_date)

    dates = []
    dates_source = []
    data = []
    ligne = ligne_depart

    while True:
        date_brute = ws.cell(ligne, col_date).value

        if not est_date_excel(date_brute):
            break

        valeurs = []

        for j in range(len(secteurs)):
            valeur = ws.cell(
                ligne,
                col_start + j,
            ).value
            valeurs.append(
                float(valeur)
                if est_nombre(valeur)
                else np.nan
            )

        date_source = convertir_date_excel(date_brute)
        dates.append(date_source)
        dates_source.append(date_source)
        data.append(valeurs)
        ligne += 1

    df = pd.DataFrame(
        data,
        index=dates,
        columns=secteurs,
    )

    cle = f"{cle_prefixe}{ws.title}_{colonne_depart}"
    return figer_historique(
        df,
        dates_source,
        cle,
    )


def lire_bloc_returns(
    ws,
    config_volatilite=None,
    secteurs=None,
    cle_prefixe="",
):
    """Lit les rendements sectoriels utilisés par Volatility."""
    cfg = config_volatilite or MODELE_EU["volatilite"]
    secteurs = secteurs or MODELE_EU["secteurs"]

    ligne = cfg["ligne_debut_returns"]
    col_date = column_index_from_string(
        cfg["colonne_date"]
    )
    col_start = column_index_from_string(
        cfg["colonne_debut_returns"]
    )

    dates = []
    dates_source = []
    data = []

    while True:
        date_brute = ws.cell(
            ligne,
            col_date,
        ).value

        if not est_date_excel(date_brute):
            break

        valeurs = []

        for j in range(len(secteurs)):
            valeur = ws.cell(
                ligne,
                col_start + j,
            ).value
            valeurs.append(
                float(valeur)
                if est_nombre(valeur)
                else np.nan
            )

        date_source = convertir_date_excel(date_brute)
        dates.append(date_source)
        dates_source.append(date_source)
        data.append(valeurs)
        ligne += 1

    df = pd.DataFrame(
        data,
        index=dates,
        columns=secteurs,
    )

    cle = f"{cle_prefixe}{ws.title}_returns"
    return figer_historique(
        df,
        dates_source,
        cle,
    )


def lire_benchmark_returns(
    ws,
    config_volatilite=None,
    cle_prefixe="",
):
    """Lit les rendements mensuels du benchmark du marché."""
    cfg = config_volatilite or MODELE_EU["volatilite"]

    ligne = cfg["ligne_debut_returns"]
    col_date = column_index_from_string(
        cfg["colonne_date"]
    )
    col_benchmark = column_index_from_string(
        cfg["colonne_benchmark_returns"]
    )

    dates = []
    dates_source = []
    data = []

    while True:
        date_brute = ws.cell(
            ligne,
            col_date,
        ).value

        if not est_date_excel(date_brute):
            break

        valeur = ws.cell(
            ligne,
            col_benchmark,
        ).value

        date_source = convertir_date_excel(date_brute)
        dates.append(date_source)
        dates_source.append(date_source)
        data.append(
            float(valeur)
            if est_nombre(valeur)
            else np.nan
        )
        ligne += 1

    df = pd.DataFrame(
        {"benchmark": data},
        index=dates,
    )

    cle = f"{cle_prefixe}{ws.title}_benchmark"
    return figer_historique(
        df,
        dates_source,
        cle,
    )


def lire_macro_externe(
    wb_macro,
    config_macro=None,
    poids_regime=None,
    cle_prefixe="",
):
    """Lit le macro score et le régime du marché."""
    cfg = config_macro or MODELE_EU["macro"]
    poids_regime = (
        poids_regime
        or MODELE_EU["poids_regime"]
    )

    ws = wb_macro[cfg["sheet"]]

    ligne = cfg["ligne_debut"]
    col_date = column_index_from_string(
        cfg["colonne_date"]
    )
    col_score = column_index_from_string(
        cfg["colonne_score"]
    )
    col_regime = column_index_from_string(
        cfg["colonne_regime"]
    )

    dates_source = []
    data = []

    while True:
        date_brute = ws.cell(
            ligne,
            col_date,
        ).value

        if not est_date_excel(date_brute):
            break

        score = ws.cell(
            ligne,
            col_score,
        ).value
        regime = ws.cell(
            ligne,
            col_regime,
        ).value
        date_source = convertir_date_excel(date_brute)

        dates_source.append(date_source)
        data.append(
            {
                "macro_score": (
                    float(score)
                    if est_nombre(score)
                    else np.nan
                ),
                "cycle": (
                    regime
                    if regime in poids_regime
                    else None
                ),
            }
        )
        ligne += 1

    df = pd.DataFrame(
        data,
        index=dates_source,
    )

    df = figer_historique(
        df,
        dates_source,
        f"{cle_prefixe}macro_outputs",
    )

    return df.sort_index()


def lire_taux_us10y(
    wb_secteur,
    config_signal=None,
    cle_prefixe="",
):
    """Lit la série US 10Y utilisée par le rate overlay."""
    cfg = config_signal or MODELE_EU["signal_taux"]
    ws = wb_secteur[cfg["sheet"]]

    ligne = cfg["ligne_debut"]
    col_date = column_index_from_string(
        cfg["colonne_date"]
    )
    col_us10y = column_index_from_string(
        cfg["colonne_us10y"]
    )

    dates_source = []
    data = []

    while True:
        date_brute = ws.cell(
            ligne,
            col_date,
        ).value

        if not est_date_excel(date_brute):
            break

        taux = ws.cell(
            ligne,
            col_us10y,
        ).value
        date_source = convertir_date_excel(date_brute)

        dates_source.append(date_source)
        data.append(
            {
                "us10y": (
                    float(taux)
                    if est_nombre(taux)
                    else np.nan
                ),
            }
        )
        ligne += 1

    df = pd.DataFrame(
        data,
        index=dates_source,
    )

    df = figer_historique(
        df,
        dates_source,
        f"{cle_prefixe}signal_taux_us10y",
    )

    return df.sort_index()
