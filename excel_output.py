"""Tableaux et présentation des classeurs sectoriels.

1. Valeurs : portefeuille, grille Output, détails et historiques.
2. Présentation : formats, liens et placement du Dashboard.
3. Écriture : remplir les classeurs avec Excel pour Windows.

Les scores sont déjà calculés par model_secto. Ce fichier les présente.
"""

from pathlib import Path
from datetime import datetime
from tempfile import TemporaryDirectory
import shutil

import numpy as np
import pandas as pd
from openpyxl.utils import column_index_from_string, get_column_letter
from local_config import PORTFOLIO_DELTA, PORTFOLIO_TILT
from model_secto import compute_deviation_from_mean
from data_utils import normalize_metadata, preserve_inputs, save_table


PILLAR_COLUMNS = {
    "EU": dict(
        zip(
            ["Leverage", "Margin", "Value", "Momentum", "Growth", "Volatility"],
            ["I", "S", "AC", "AM", "AU", "BE"],
        )
    ),
    "US": dict(
        zip(
            ["Leverage", "Margin", "Value", "Momentum", "Growth", "Volatility"],
            ["I", "P", "Z", "AJ", "AR", "BE"],
        )
    ),
}

# Score, mesure brute, historique.
DETAIL_COLUMNS = {
    "EU": {
        "net_debt_ebitda": ("J", "L", "K"),
        "fcf_total_debt": ("M", "O", "N"),
        "debt_equity": ("P", "R", "Q"),
        "operating_margin": ("T", "V", "U"),
        "net_margin": ("W", "Y", "X"),
        "ebitda_margin": ("Z", "AB", "AA"),
        "price_fcf": ("AD", "AF", "AE"),
        "ev_ebitda": ("AG", "AI", "AH"),
        "price_sales": ("AJ", "AL", "AK"),
        "momentum_6m_1m": ("AN", "AO", None),
        "momentum_12m_1m": ("AP", "AQ", None),
        "earnings_revision_ratio": ("AR", "AT", "AS"),
        "eps_growth": ("AV", "AX", "AW"),
        "sales_growth": ("AY", "BA", "AZ"),
        "ebitda_growth": ("BB", "BD", "BC"),
        "volatility_6m": ("BF", "BG", None),
        "downside_volatility_18m": ("BH", "BI", None),
    },
    "US": {
        "net_debt_ebitda": ("J", "L", "K"),
        "fcf_total_debt": ("M", "O", "N"),
        "operating_margin": ("Q", "S", "R"),
        "net_margin": ("T", "V", "U"),
        "ebitda_margin": ("W", "Y", "X"),
        "price_fcf": ("AA", "AC", "AB"),
        "ev_ebitda": ("AD", "AF", "AE"),
        "price_sales": ("AG", "AI", "AH"),
        "momentum_6m_1m": ("AK", "AL", None),
        "momentum_12m_1m": ("AM", "AN", None),
        "earnings_revision_ratio": ("AO", "AQ", "AP"),
        "eps_growth": ("AS", "AU", "AT"),
        "sales_growth": ("AV", "AX", "AW"),
        "ebitda_growth": ("AY", "BA", "AZ"),
        "cash_flow_growth": ("BB", "BD", "BC"),
        "volatility_6m": ("BF", "BG", None),
        "downside_volatility_18m": ("BH", "BI", None),
    },
}


# 1. Valeurs à écrire dans Output


def compute_portfolio(results, end_date):
    """Applique les recommandations du mois précédent aux poids du benchmark."""
    history = results["history"]
    valid_dates = history.loc[history["score_global_macro"].notna(), "date"].unique()
    recommendation_history = (
        history[history["date"].isin(valid_dates)]
        .pivot(index="date", columns="secteur", values="recommendation")
        .sort_index()
    )
    weight_history = results["benchmark_weights"].sort_index()
    sector_return_history = results["returns"].sort_index()
    benchmark_returns = results["benchmark_returns"]["benchmark"]
    rows = []
    for date, sector_returns in sector_return_history.loc[:end_date].iterrows():
        signal_date = date - pd.offsets.MonthEnd(1)
        if signal_date not in recommendation_history.index:
            continue
        available_weights = weight_history.loc[:date]
        if available_weights.empty:
            continue
        recommendations = (
            recommendation_history.loc[signal_date].reindex(sector_returns.index).fillna("Neutral")
        )
        benchmark_weights = available_weights.iloc[-1].reindex(sector_returns.index)
        adjusted_weights = benchmark_weights.copy()
        positive, negative = recommendations.eq("Positive"), recommendations.eq("Negative")
        adjusted_weights[positive] = np.maximum(
            benchmark_weights[positive] * (1 + PORTFOLIO_TILT), benchmark_weights[positive] + PORTFOLIO_DELTA
        )
        adjusted_weights[negative] = np.minimum(
            benchmark_weights[negative] * (1 - PORTFOLIO_TILT),
            (benchmark_weights[negative] - PORTFOLIO_DELTA).clip(lower=0),
        )
        if benchmark_weights.isna().any() or sector_returns[adjusted_weights > 0].isna().any():
            portfolio = np.nan
        else:
            portfolio = (adjusted_weights * sector_returns).sum() / adjusted_weights.sum()
        rows.append(
            [
                date,
                benchmark_returns.get(date, np.nan),
                portfolio,
                sector_returns[positive].mean(),
                sector_returns[negative].mean(),
            ]
        )
    monthly = pd.DataFrame(
        rows, columns=["date", "benchmark", "portfolio", "top", "bottom"]
    ).set_index("date")
    if monthly.empty:
        raise ValueError("Aucun rendement de portefeuille disponible.")
    monthly["long_short"] = monthly["top"] - monthly["bottom"]
    # Un mois manquant interrompt le cumul plutôt que de créer une performance.
    cumulative = (1 + monthly).cumprod(skipna=False) * 100
    start = monthly.index[0] - pd.offsets.MonthEnd(1)
    cumulative.loc[start] = 100.0
    return monthly, cumulative.sort_index()


def output_layout(config):
    """Place les blocs avec deux lignes libres entre eux."""
    n = len(config["sectors"])
    sparkline_count = (
        sum(spark is not None for _, _, spark in DETAIL_COLUMNS[config["name"]].values())
        * n
    )
    measure_count = sum(len(v) for v in config["backtest_variables"].values()) * n
    recommendations = 22
    history = recommendations + n + 3
    details = history + sparkline_count + 3
    curves = details + measure_count + 3
    return dict(recommendations=recommendations, history=history, details=details, curves=curves)


def output_cell(address, market):
    """Relie une cellule du scorecard à sa nouvelle position dans Output."""
    headers = {
        "BO4": "B2",
        "BP4": "B9",
        "BQ4": "B10",
        "BP7": "B11",
        "BV7": "B12",
        "BP6": "B11",
        "BV6": "B12",
        "BJ8": "B13",
        "BK10": "B6",
    }
    if address in headers:
        return headers[address]
    column = "".join(c for c in address if c.isalpha())
    row_number = int("".join(c for c in address if c.isdigit()))
    if column == "CD":
        start = 9 if market == "EU" else 10
        return f"B{15 + row_number - start}"
    return f"{get_column_letter(column_index_from_string(column) - 4)}{row_number - 8}"


def _set_output_value(grid, address, value):
    column = "".join(char for char in address if char.isalpha())
    row_number = int("".join(char for char in address if char.isdigit()))
    grid.iat[row_number - 1, column_index_from_string(column) - 1] = (
        "n.a." if pd.isna(value) else value
    )


def _write_output_block(grid, row_number, column, values):
    values = pd.DataFrame(values).astype(object).where(pd.notna(values), None)
    start_row, start_column = row_number - 1, column - 1
    grid.iloc[
        start_row:start_row + len(values),
        start_column:start_column + values.shape[1],
    ] = values.to_numpy()


def _month_label(date):
    months = ['Jan', 'Fev', 'Mar', 'Avr', 'Mai', 'Juin', 'Juil', 'Aout', 'Sept', 'Oct', 'Nov', 'Dec']
    return f"{months[date.month - 1]}-{date.year}"


def _write_scorecard(grid, results, config, date):
    """Remplit les scores du mois et leur comparaison avec le mois précédent."""
    sectors = config["sectors"]
    market = config["name"]
    history = results["history"]
    previous = date - pd.offsets.MonthEnd(1)
    current = history[history["date"].eq(date)].set_index("secteur")
    previous_data = history[history["date"].eq(previous)].set_index("secteur")

    def scorecard(address, value):
        _set_output_value(grid, output_cell(address, market), value)

    returns = results["returns"]
    bench = results["benchmark_returns"]["benchmark"]
    recommendation_codes = {"Negative": -1, "Neutral": 0, "Positive": 1}
    details = DETAIL_COLUMNS[market]
    scorecard("H9", "Secteur")
    for pillar, column in PILLAR_COLUMNS[market].items():
        scorecard(column + "9", pillar)
    for variable, (score, raw, _) in details.items():
        scorecard(score + "9", "Score " + variable.replace("_", " "))
        scorecard(raw + "9", variable.replace("_", " "))
    for col, title in {
        "BJ": "Macro",
        "BM": "Secteur",
        "BN": "Votes Top",
        "BO": "Votes Worst",
        "BP": "Score global",
        "BQ": "Reco",
        "BR": "Variation reco",
        "BU": "Votes Top précédents",
        "BV": "Votes Worst précédents",
        "BW": "Score précédent",
        "BX": "Reco précédente",
        "BY": _month_label(date) + " EQ Performance",
        "BZ": _month_label(date)
        + (
            " XSRet vs FTSE World Europe"
            if market == "EU"
            else " XSRet vs " + config["volatility"]["benchmark_name"]
        ),
    }.items():
        scorecard(col + "9", title)
    for i, sector in enumerate(sectors):
        row_number = i + 10
        scorecard(f"H{row_number}", sector)
        scorecard(f"BM{row_number}", sector)
        for pillar, column in PILLAR_COLUMNS[market].items():
            rank = results["ranks"][pillar].at[date, sector]
            scorecard(f"{column}{row_number}", len(sectors) + 1 - rank)
        for variable, (score_col, raw_column, _) in details.items():
            score = results["variable_scores"].get(variable)
            raw = results["variable_raw"].get(variable)
            scorecard(
                f"{score_col}{row_number}",
                score.at[date, sector] if score is not None else np.nan,
            )
            scorecard(f"{raw_column}{row_number}", raw.at[date, sector] if raw is not None else np.nan)
        row = current.loc[sector]
        scorecard(f"BJ{row_number}", len(sectors) + 1 - row["rang_tilt_macro"])
        for column, field in [
            ("BN", "top_count"),
            ("BO", "bottom_count"),
            ("BP", "score_global_macro"),
            ("BQ", "recommendation"),
        ]:
            scorecard(f"{column}{row_number}", row[field])
        for column, field in [
            ("BU", "top_count"),
            ("BV", "bottom_count"),
            ("BW", "score_global_macro"),
            ("BX", "recommendation"),
        ]:
            scorecard(
                f"{column}{row_number}",
                (
                    previous_data.at[sector, field]
                    if sector in previous_data.index
                    else np.nan
                ),
            )
        previous_recommendation = (
            previous_data.at[sector, "recommendation"]
            if sector in previous_data.index
            else None
        )
        scorecard(
            f"BR{row_number}",
            (
                recommendation_codes[row["recommendation"]]
                - recommendation_codes[previous_recommendation]
                if previous_recommendation
                else np.nan
            ),
        )
        ret = returns.at[date, sector]
        scorecard(f"BY{row_number}", ret)
        scorecard(f"BZ{row_number}", ret - bench.at[date])



def _write_indicator_history(grid, results, config, date, layout):
    """Écrit les mini-historiques puis les mesures détaillées, sans modifier les résultats."""
    sectors = config["sectors"]
    details = DETAIL_COLUMNS[config["name"]]
    # Chaque colonne représente un mois, du plus récent au plus ancien.
    history_dates = pd.date_range(end=date, periods=60, freq="ME")[::-1]
    _write_output_block(grid, layout["history"], 4, [["Variable", "Secteur"] + history_dates.tolist()])
    history_row = layout["history"] + 1
    for variable, (_, _, spark_col) in details.items():
        if spark_col is None:
            continue
        raw = results["variable_raw"].get(variable)
        if raw is None:
            raw = pd.DataFrame(index=history_dates, columns=sectors, dtype=float)
        definition = next(
            (
                variables[variable]
                for variables in config["historical_variables"].values()
                if variable in variables
            ),
            None,
        )
        deviation = (
            compute_deviation_from_mean(raw, definition["exclude_financials_from_mean"])
            if definition
            else raw
        )
        deviation = deviation.reindex(history_dates)
        for sector in sectors:
            _write_output_block(grid, history_row, 4, [[variable, sector] + deviation[sector].tolist()])
            history_row += 1

    _write_output_block(grid, layout["details"], 4, [["Date", "Secteur", "Pilier", "Variable", "Mesure", "Score"]])
    rows = []
    for pillar, variables in config["backtest_variables"].items():
        for variable in variables:
            raw = results["variable_raw"].get(variable)
            score = results["variable_scores"].get(variable)
            for sector in sectors:
                rows.append(
                    [
                        date,
                        sector,
                        pillar,
                        variable,
                        raw.at[date, sector] if raw is not None else None,
                        score.at[date, sector] if score is not None else None,
                    ]
                )
    _write_output_block(grid, layout["details"] + 1, 4, rows)



def build_output(results, config):
    """Prépare les scores, les historiques et les séries des graphiques."""
    sectors = config["sectors"]
    market = config["name"]
    history = results["history"]
    if history.empty:
        raise ValueError("Aucune recommandation complète disponible.")
    date = history["date"].max()
    current = history[history["date"].eq(date)].set_index("secteur")
    monthly, cumulative = compute_portfolio(results, date)
    layout = output_layout(config)
    grid = pd.DataFrame(
        None,
        index=range(layout["curves"] + len(cumulative)),
        columns=range(74),
        dtype=object,
    )

    pillar_date = max(df.index.max() for df in results["pillars"].values())
    cycle = current["cycle"].iloc[0]
    rate_signal = current["signal_taux"].iloc[0]
    regime = {"C": "Contraction", "R": "Recovery", "E": "Expansion", "SD": "Slowdown"}[
        cycle
    ]
    next_month = date + pd.offsets.MonthBegin(1)
    _write_output_block(grid,
        1,
        1,
        [
            ["Marché", market],
            ["Date modèle", date],
            ["Derniers piliers", pillar_date],
            ["Allocation", next_month],
            ["Cycle", cycle],
            ["Signal taux", rate_signal],
            ["Derniers poids benchmark", results["benchmark_weights"].index.max()],
            ["Derniers rendements", results["returns"].index.max()],
            ["Mois allocation", _month_label(next_month)],
            ["Mois signal", _month_label(date)],
            ["Classement courant", "Sectors Ranking " + _month_label(next_month)],
            ["Classement précédent", "Sectors Ranking " + _month_label(date)],
            ["Classement macro", "Macro rank - Currently in " + regime],
            ["Référence", "Bench"],
            ["Top", monthly.at[date, "top"]],
            ["Short", -monthly.at[date, "bottom"]],
            ["Long / Short", monthly.at[date, "long_short"]],
            ["Portefeuille", monthly.at[date, "portfolio"]],
            ["Benchmark", monthly.at[date, "benchmark"]],
        ],
    )

    _write_scorecard(grid, results, config, date)

    returns = results["returns"]
    recommendation_codes = {"Negative": -1, "Neutral": 0, "Positive": 1}
    recommendation_dates = pd.date_range(end=date, periods=15, freq="ME")
    recommendations = history.pivot(index="secteur", columns="date", values="recommendation")
    _write_output_block(grid,
        layout["recommendations"],
        4,
        [[market + " Reco"] + [d + pd.offsets.MonthBegin(1) for d in recommendation_dates]],
    )
    for i, sector in enumerate(sectors):
        _write_output_block(grid,
            layout["recommendations"] + 1 + i,
            4,
            [
                [sector]
                + [
                    (
                        recommendation_codes.get(recommendations.at[sector, d], None)
                        if d in recommendations.columns
                        else None
                    )
                    for d in recommendation_dates
                ]
            ],
        )

    title = ("Europe" if market == "EU" else "US") + " equity sectors Performance"
    _set_output_value(grid, f"W{layout['recommendations']}", title)
    _write_output_block(grid, layout["recommendations"] + 1, 22, [[""] + sectors])
    series = returns.loc[returns.index <= date].sort_index()
    horizons = [("1M", 1), ("3M", 3), ("YTD", date.month), ("1Y", 12), ("3Y", 36), ("5Y", 60)]
    for i, (name, period_count) in enumerate(horizons):
        perf = (1 + series.tail(period_count)).prod(min_count=period_count) - 1
        _write_output_block(grid,
            layout["recommendations"] + 2 + i,
            22,
            [[name] + [v if pd.notna(v) else "n.a." for v in perf]],
        )

    _write_indicator_history(grid, results, config, date, layout)

    titles = [
        "Date",
        "CumRet_" + config["volatility"]["benchmark_name"],
        "CumRet_Sector deviations portfolio",
        "Top 3 Sectors" if market == "EU" else "CumRet_Top",
        "Worst 3 sectors" if market == "EU" else "CumRet_Bottom",
        "CumRet_Long Short",
    ]
    _write_output_block(grid, layout["curves"], 4, [titles])
    _write_output_block(grid, layout["curves"] + 1, 4, cumulative.reset_index().values.tolist())
    one_year = cumulative.loc[
        cumulative.index >= date - pd.offsets.MonthEnd(12), ["benchmark", "portfolio", "top", "bottom"],
    ]
    one_year = one_year / one_year.iloc[0] * 100
    _write_output_block(grid, layout["curves"], 13, [[titles[0], titles[1], titles[2], titles[3], titles[4]]])
    _write_output_block(grid, layout["curves"] + 1, 13, one_year.reset_index().values.tolist())
    return grid, monthly, cumulative


def save_excel(results, output_dir, config):
    """Exporte la feuille de valeurs et les rendements mensuels."""
    directory = Path(output_dir)
    directory.mkdir(parents=True, exist_ok=True)
    grid, monthly, _ = build_output(results, config)
    file_path = directory / f"{config['output_prefix']}_output.xlsx"
    grid.to_excel(file_path, sheet_name="Output", index=False, header=False)
    normalize_metadata(file_path)
    save_table(monthly.reset_index(), directory / f"{config['output_prefix']}_performance_portefeuille.xlsx")
    return file_path


# 2. Présentation du classeur


def excel_value(value):
    if pd.isna(value):
        return None
    if isinstance(value, (pd.Timestamp, datetime)):
        return (pd.Timestamp(value) - pd.Timestamp("1899-12-30")).total_seconds() / 86400
    return value


def append_recommendations(sheet, history, sectors):
    """Ajoute les nouveaux mois sans modifier les recommandations archivées."""
    last_row = sheet.Cells(sheet.Rows.Count, 1).End(-4162).Row
    last_date = pd.Timestamp("1899-12-30") + pd.Timedelta(days=sheet.Cells(last_row, 1).Value2)
    new_recommendations = history[history["date"] > last_date]
    rows = []
    for date, group in new_recommendations.groupby("date", sort=True):
        recommendations = group.set_index("secteur")["recommendation"].reindex(sectors)
        if recommendations.isna().any():
            raise ValueError(f"Recommandations incomplètes : {date:%d/%m/%Y}.")
        rows.append(
            [excel_value(date)]
            + recommendations.eq("Positive").astype(int).tolist()
            + recommendations.eq("Negative").astype(int).tolist()
        )
    if rows:
        width = 1 + 2 * len(sectors)
        end = get_column_letter(width)
        destination = sheet.Range(f"A{last_row+1}:{end}{last_row+len(rows)}")
        sheet.Range(f"A{last_row}:{end}{last_row}").Copy()
        destination.PasteSpecial(Paste=-4122)
        destination.Value2 = tuple(tuple(row_number) for row_number in rows)
    return len(rows)


def fill_workbook(workbook, results, config, export_file):
    """Remplace les valeurs en conservant la présentation du modèle Excel."""
    workbook.Activate()
    grid, _, _ = build_output(results, config)
    sheet = workbook.Worksheets("Output")
    last_row = sheet.UsedRange.Rows.Count
    last_column = get_column_letter(grid.shape[1])

    # 1. Conserver les formats du modèle ; effacer seulement les anciennes valeurs.
    sheet.UsedRange.ClearContents()
    if len(grid) > last_row:
        # Prolonger le format de la dernière ligne pour les nouveaux mois.
        sheet.Range(f"A{last_row}:{last_column}{last_row}").Copy()
        sheet.Range(f"A{last_row+1}:{last_column}{len(grid)}").PasteSpecial(Paste=-4122)
        workbook.Application.CutCopyMode = False

    values = tuple(tuple(excel_value(v) for v in row) for row in grid.to_numpy())
    sheet.Range(f"A1:{last_column}{len(grid)}").Value2 = values
    if (
        sheet.Range("B2").Value2 != excel_value(results["history"]["date"].max())
        or sheet.Range("D1").Value2 != "Secteur"
    ):
        raise RuntimeError("Les valeurs n'ont pas été écrites dans Output.")

    # 2. Ajouter les recommandations nouvelles sans toucher à l'historique.
    added_months = append_recommendations(
        workbook.Worksheets("Reco_histo"), results["history"], config["sectors"]
    )

    # 3. Actualiser les graphiques en réutilisant les liens du modèle.
    dashboard = workbook.Worksheets("Dashboard")
    sheet.Calculate()
    dashboard.Calculate()
    for chart in dashboard.ChartObjects():
        chart.Chart.Refresh()

    # 4. Copier les valeurs et les formats dans le fichier Output autonome.
    export = workbook.Application.Workbooks.Add()
    try:
        export_sheet = export.Worksheets(1)
        export_sheet.Name = "Output"
        source = sheet.Range(f"A1:{last_column}{len(grid)}")
        source.Copy(Destination=export_sheet.Range("A1"))
        source.Copy()
        export_sheet.Range("A1").PasteSpecial(Paste=8)  # Largeurs des colonnes.
        workbook.Application.CutCopyMode = False
        for row_number in range(1, len(grid) + 1):
            export_sheet.Rows(row_number).RowHeight = sheet.Rows(row_number).RowHeight
        export.SaveAs(str(export_file), FileFormat=51)
    finally:
        export.Close(False)

    # 5. Enregistrer le modèle actualisé sur la page Dashboard.
    workbook.Activate()
    dashboard.Activate()
    workbook.Save()
    return added_months


# 3. Mise à jour des deux classeurs avec Excel pour Windows


def update_workbooks(models, output_dir):
    """Prépare les deux fichiers, puis remplace les originaux après succès."""
    import win32com.client

    added_months = {}  # Nombre de mois ajoutés à Reco_histo pour chaque marché.
    with TemporaryDirectory(prefix="sector_excel_") as temporary:  # Dossier de travail provisoire.
        work_dir = Path(temporary)
        app = win32com.client.gencache.EnsureDispatch(
            win32com.client.DispatchEx("Excel.Application")
        )  # Nouvelle instance Excel, indépendante des classeurs ouverts.
        app.Visible = True  # Affiche Excel pendant la mise à jour.
        app.DisplayAlerts = False  # Évite les fenêtres de confirmation.
        app.EnableEvents = False  # Désactive les événements Excel pendant la mise à jour.
        app.AskToUpdateLinks = False  # N'interroge pas les sources liées.
        blank_workbook = app.Workbooks.Add()  # Classeur de support pour la session Excel.
        app.Calculation = -4135  # -4135 = calcul manuel, sans recalcul global des entrées.
        app.CalculateBeforeSave = False  # N'impose pas un recalcul lors de l'enregistrement.

        try:
            for model in models:  # Traite EU et US avec les mêmes étapes.
                print(f"Mise à jour Excel {model.universe}...", flush=True)
                original = model.excel_file  # Classeur du marché utilisé comme modèle.
                copy_path = work_dir / original.name  # Fichier modifié avant validation.
                shutil.copy2(original, copy_path)  # L'original reste intact pendant le calcul.
                wb = app.Workbooks.Open(
                    str(copy_path), UpdateLinks=0, ReadOnly=False
                )  # 0 : liens externes non actualisés ; False : ouverture en écriture.
                export = output_dir / f"{model.config['output_prefix']}_output.xlsx"
                try:
                    if wb.ReadOnly:  # Vérifie que la copie est réellement modifiable.
                        raise RuntimeError(f"Classeur en lecture seule : {copy_path.name}")
                    added_months[model.universe] = fill_workbook(
                        wb, model.results, model.config, export
                    )  # Écrit Output, complète Reco_histo et rafraîchit Dashboard.
                finally:
                    wb.Close(False)  # Ferme la copie, déjà enregistrée par fill_workbook.

                preserve_inputs(original, copy_path)  # Garde les entrées et formules du modèle.
                normalize_metadata(copy_path)  # Normalise les métadonnées du classeur.
                normalize_metadata(export)  # Normalise celles de l'export Output.
        finally:
            blank_workbook.Close(False)  # Ferme le classeur de support sans sauvegarde.
            app.Quit()  # Arrête uniquement cette instance Excel.

        # Remplacer les originaux seulement après la réussite des deux marchés.
        for model in models:
            shutil.copy2(work_dir / model.excel_file.name, model.excel_file)
    return added_months
