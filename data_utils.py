# -*- coding: utf-8 -*-
"""Utilitaires de données et de fichiers Excel.

1. Lecture des fichiers et dates mensuelles.
2. Historique : conserver la première observation de chaque mois.
3. Lecture des blocs, rendements, macro et taux.
4. Conservation des entrées, des graphiques et des métadonnées Excel.

Les chemins et les paramètres sont définis dans local_config.py.
"""

from copy import deepcopy
from pathlib import Path
import re
from zipfile import ZipFile, ZIP_DEFLATED

import numpy as np
import pandas as pd
from lxml import etree as ET
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from local_config import HISTORY_CONFIG, MACRO_EXCEL_FILE, EU_CONFIG, US_CONFIG

NS = {
    "s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}
S = "{" + NS["s"] + "}"
HISTORY_SHEET = "Historique_data"
_history = None
_history_changed = False
_history_file = None

# 1. Fichiers et dates


def open_workbooks(excel_file=None, macro_file=None):
    """Lit les valeurs enregistrées des deux fichiers source."""
    excel_file = Path(excel_file or EU_CONFIG["excel_file"])
    macro_file = Path(macro_file or MACRO_EXCEL_FILE)
    sector_workbook = load_workbook(excel_file, data_only=True)
    macro_workbook = load_workbook(macro_file, data_only=True)
    return sector_workbook, macro_workbook


def is_number(x):
    """Indique si une cellule contient une valeur numérique exploitable."""
    return isinstance(x, (int, float, np.integer, np.floating)) and not pd.isna(x)


def convert_excel_date(x):
    """Convertit une date Excel ou Python en Timestamp pandas."""
    if isinstance(x, (pd.Timestamp, np.datetime64)):
        return pd.Timestamp(x)

    if hasattr(x, "year") and hasattr(x, "month") and hasattr(x, "day"):
        return pd.Timestamp(x)

    if is_number(x):
        return pd.Timestamp("1899-12-30") + pd.to_timedelta(float(x), unit="D")

    return pd.NaT


def is_excel_date(x):
    """Vérifie qu'une cellule contient une date exploitable."""
    date = convert_excel_date(x)

    if pd.isna(date):
        return False

    return 1980 <= date.year <= 2100


def normalize_month_end(x):
    """Utilise le dernier jour calendaire du mois comme clé de date."""
    date = convert_excel_date(x)

    if pd.isna(date):
        return pd.NaT

    return pd.Timestamp(date).to_period("M").to_timestamp("M")


def normalize_monthly_index(df):
    """Normalise l'index en fins de mois et conserve une ligne par mois."""
    result = df.copy()
    result.index = pd.DatetimeIndex([normalize_month_end(x) for x in result.index])
    result = result[~result.index.duplicated(keep="first")]
    return result.sort_index(ascending=False)


# 2. Historique mensuel


def history_path():
    """Retourne le fichier Excel qui contient tout l'historique."""
    return Path(__file__).resolve().parent / HISTORY_CONFIG["file"]


def load_history(reload=False, file_path=None, fresh=False):
    """Charge le snapshot choisi, ou initialise un historique vide pour rebuild."""
    global _history, _history_changed, _history_file
    if _history is None or reload or fresh:
        _history_file = Path(file_path) if file_path is not None else history_path()
        _history = (
            pd.DataFrame() if fresh or not _history_file.exists()
            else pd.read_excel(_history_file, sheet_name=HISTORY_SHEET)
        )
        for column in ("date", "date_source"):
            if column in _history:
                _history[column] = pd.to_datetime(_history[column])
        _history_changed = False
    return _history


def save_history():
    """Écrit les nouveaux mois sans remplacer les observations figées."""
    global _history_changed
    if not _history_changed:
        return
    file_path = _history_file or history_path()
    temporary = file_path.with_name(file_path.stem + ".tmp.xlsx")
    _history.sort_values(["series_key", "date"], ascending=[True, False]).to_excel(
        temporary, sheet_name=HISTORY_SHEET, index=False, freeze_panes=(1, 3)
    )
    normalize_metadata(temporary)
    temporary.replace(file_path)
    _history_changed = False


def freeze_history(df, source_dates, key):
    """Conserve la première ligne mensuelle d'une série, y compris ses valeurs absentes.

    L'entrée et la sortie ont les mois en index et les secteurs en colonnes.
    Seul l'archivage passe au format long (series_key, date, date_source).
    """
    global _history, _history_changed
    monthly = normalize_monthly_index(df)
    source_dates = pd.Series(
        [convert_excel_date(date) for date in source_dates],
        index=[normalize_month_end(date) for date in source_dates],
        name="date_source",
    )
    source_dates = source_dates[~source_dates.index.duplicated(keep="first")]

    observations = monthly.copy()
    observations.insert(0, "date_source", source_dates.reindex(observations.index).values)
    observations.insert(0, "date", observations.index)
    observations.insert(0, "series_key", key)
    observations = observations.reset_index(drop=True)
    columns = list(df.columns)
    observations = observations[observations[columns].notna().any(axis=1)]

    if not HISTORY_CONFIG.get("enabled", True):
        return monthly

    history = load_history()
    existing_dates = (
        history.loc[history["series_key"] == key, "date"]
        if not history.empty else pd.Series(dtype="datetime64[ns]")
    )
    new_months = observations.loc[~observations["date"].isin(existing_dates)]
    if not new_months.empty:
        history = pd.concat([history, new_months], ignore_index=True, sort=False)
        _history = history
        _history_changed = True
    if history.empty:
        return monthly

    archived = history[history["series_key"] == key].copy()
    archived = archived.sort_values("date", ascending=False).set_index("date")
    for column in columns:
        if column not in archived.columns:
            archived[column] = np.nan
    return archived[columns].copy()


def read_history_series(key, columns):
    """Restitue une série mensuelle enregistrée, sans lire les fichiers source."""
    history = load_history()
    series = history.loc[history["series_key"].eq(key)].copy() if not history.empty else pd.DataFrame()
    if series.empty:
        raise ValueError(f"Série absente de l'historique : {key}")
    missing = [column for column in columns if column not in series]
    if missing:
        raise ValueError(f"Colonnes absentes de {key} : {', '.join(missing)}")
    return series.sort_values("date", ascending=False).set_index("date")[columns].copy()


# 3. Blocs de données


def read_sector_block(
    ws, start_column, sectors=None, key_prefix="", start_row=8, date_column="A",
    sheet_name=None,
):
    """Lit le bloc Excel ou la série figée correspondante."""
    sectors = sectors or EU_CONFIG["sectors"]
    if ws is None:
        return read_history_series(f"{key_prefix}{sheet_name}_{start_column}", sectors)
    first_column = column_index_from_string(start_column)
    date_column_number = column_index_from_string(date_column)
    source_dates, rows = [], []
    row_number = start_row

    while True:
        raw_date = ws.cell(row_number, date_column_number).value
        if not is_excel_date(raw_date):
            break
        values = []
        for offset in range(len(sectors)):
            value = ws.cell(row_number, first_column + offset).value
            values.append(float(value) if is_number(value) else np.nan)
        source_dates.append(convert_excel_date(raw_date))
        rows.append(values)
        row_number += 1

    raw = pd.DataFrame(rows, index=source_dates, columns=sectors)
    key = f"{key_prefix}{ws.title}_{start_column}"
    return freeze_history(raw, source_dates, key)


def read_return_indices(ws, sectors):
    """Lit les niveaux d'indice : benchmark puis secteurs, dates décroissantes."""
    dates, rows = [], []
    for row_number in range(7, ws.max_row + 1):
        date = ws.cell(row_number, 1).value
        if not is_excel_date(date):
            break
        values = []
        for column in range(2, len(sectors) + 3):
            value = ws.cell(row_number, column).value
            values.append(float(value) if is_number(value) else np.nan)
        dates.append(convert_excel_date(date))
        rows.append(values)

    indices = pd.DataFrame(rows, index=dates, columns=["benchmark"] + sectors)
    return indices.sort_index(ascending=False)


def read_sector_returns(ws, volatility_config=None, sectors=None, key_prefix="", sheet_name=None):
    """Retourne les rendements sectoriels depuis Excel ou le snapshot figé."""
    sectors = sectors or EU_CONFIG["sectors"]
    if ws is None:
        return read_history_series(f"{key_prefix}{sheet_name}_returns", sectors)
    indices = read_return_indices(ws, sectors)
    returns = indices / indices.shift(-1) - 1
    return freeze_history(returns[sectors], returns.index, f"{key_prefix}{ws.title}_returns")


def read_benchmark_returns(ws, volatility_config=None, key_prefix="", sheet_name=None):
    """Retourne le rendement benchmark depuis Excel ou le snapshot figé."""
    if ws is None:
        return read_history_series(f"{key_prefix}{sheet_name}_benchmark", ["benchmark"])
    sectors = US_CONFIG["sectors"] if key_prefix == "US_" else EU_CONFIG["sectors"]
    indices = read_return_indices(ws, sectors)
    returns = indices[["benchmark"]] / indices[["benchmark"]].shift(-1) - 1
    return freeze_history(returns, returns.index, f"{key_prefix}{ws.title}_benchmark")


def extend_macro(ws, cfg, regimes):
    """Prolonge les scores dans le classeur en mémoire ; aucun fichier n'est enregistré."""
    score_col = column_index_from_string(cfg["score_column"])
    cycle_col = column_index_from_string(cfg["regime_column"])
    rows = [
        r
        for r in range(cfg["start_row"], ws.max_row + 1)
        if is_number(ws.cell(r, score_col).value)
        and ws.cell(r, cycle_col).value in regimes
    ]
    if not rows:
        return

    def read_data(col, r):
        return ws[f"{col}{r}"].value

    def cycle(score, previous):
        if score > previous:
            return "E" if score >= 0 else "R"
        return "SD" if score >= 0 else "C"

    for r in range(rows[-1] + 1, ws.max_row + 1):
        if not is_excel_date(read_data("A", r)):
            break
        if cfg["sheet"] == "Europe":
            values = [read_data("D", r - 1), read_data("I", r), read_data("N", r)]
            if not all(is_number(v) for v in values):
                break
            score = sum(values) / 3
            previous = read_data("R", r - 1)
            simple = cycle(score, previous)
            threshold = ws["T1"].value
            max_threshold = ws["W1"].value
            filtered = (
                simple
                if (abs(score - previous) > threshold or score * previous < 0)
                else read_data("T", r - 1)
            )
            pivot = (
                score if filtered != read_data("S", r - 1) else read_data("U", r - 1)
            )
            previous_pivot = read_data("U", r - 1)
            maximum = max(read_data("R", i) for i in range(r - 12, r))
            regime = (
                simple
                if (
                    abs(pivot - previous_pivot) > threshold
                    or pivot * previous_pivot < 0
                    or abs(score - maximum) > max_threshold
                )
                else read_data("V", r - 1)
            )
            for col, value in zip(["R", "S", "T", "U", "V"], [score, simple, filtered, pivot, regime]):
                ws[f"{col}{r}"] = value
        else:
            values = [read_data("M", r), read_data("J", r), read_data("G", r - 1), read_data("D", r)]
            if not all(is_number(v) for v in values):
                break
            score = sum(v * weights for v, weights in zip(values, [0.3, 0.3, 0.3, 0.1]))
            ws[f"Y{r}"] = score
            ws[f"Z{r}"] = cycle(score, read_data("Y", r - 1))


def read_macro_data(macro_workbook, macro_config=None, regime_weights=None, key_prefix=""):
    """Lit score et cycle, puis renvoie les mois par ordre chronologique."""
    cfg = macro_config or EU_CONFIG["macro"]
    regime_weights = regime_weights or EU_CONFIG["regime_weights"]
    key = f"{key_prefix}macro_outputs_{cfg['score_column']}_{cfg['regime_column']}"
    if macro_workbook is None:
        return read_history_series(key, ["macro_score", "cycle"]).sort_index()
    sheet = macro_workbook[cfg["sheet"]]

    # Les signaux du fichier macro sont calculés et enregistrés dans Excel.
    date_column = column_index_from_string(cfg["date_column"])
    score_column = column_index_from_string(cfg["score_column"])
    regime_column = column_index_from_string(cfg["regime_column"])
    source_dates, rows = [], []
    row_number = cfg["start_row"]
    while True:
        raw_date = sheet.cell(row_number, date_column).value
        if not is_excel_date(raw_date):
            break
        score = sheet.cell(row_number, score_column).value
        regime = sheet.cell(row_number, regime_column).value
        source_dates.append(convert_excel_date(raw_date))
        rows.append({
            "macro_score": float(score) if is_number(score) else np.nan,
            "cycle": regime if regime in regime_weights else None,
        })
        row_number += 1

    macro = pd.DataFrame(rows, index=source_dates)
    # Deux couples de colonnes sont deux séries distinctes dans l'historique.
    archived = freeze_history(macro, source_dates, key)
    return archived.sort_index()


def read_us10y(sector_workbook, signal_config=None, key_prefix=""):
    """Lit le taux US 10 ans et renvoie son historique par date croissante."""
    cfg = signal_config or EU_CONFIG["rate_signal"]
    key = f"{key_prefix}signal_taux_us10y"
    if sector_workbook is None:
        return read_history_series(key, ["us10y"]).sort_index()
    sheet = sector_workbook[cfg["sheet"]]
    date_column = column_index_from_string(cfg["date_column"])
    rate_column = column_index_from_string(cfg["us10y_column"])
    source_dates, rows = [], []
    row_number = cfg["start_row"]

    while True:
        raw_date = sheet.cell(row_number, date_column).value
        if not is_excel_date(raw_date):
            break
        rate = sheet.cell(row_number, rate_column).value
        source_dates.append(convert_excel_date(raw_date))
        rows.append({"us10y": float(rate) if is_number(rate) else np.nan})
        row_number += 1

    rates = pd.DataFrame(rows, index=source_dates)
    archived = freeze_history(rates, source_dates, key)
    return archived.sort_index()


# 4. Format natif des classeurs


def normalize_metadata(file_path):
    """Conserve les noms usuels des styles et retire le chemin local."""
    from openpyxl.styles.builtins import styles

    names = {str(style.builtinId): name for name, style in styles.items()}
    font_names = {"Hans": "DengXian", "Hant": "PMingLiU", "Jpan": "Yu Gothic"}
    character_pattern = re.compile(r"[\u3400-\u9fff\uf900-\ufaff]")
    changes = {}
    with ZipFile(file_path) as archive:
        for path in archive.namelist():
            if path not in (
                "xl/workbook.xml",
                "xl/styles.xml",
                "docProps/app.xml",
            ) and not path.startswith("xl/theme/"):
                continue
            content = archive.read(path)
            document = ET.fromstring(content)
            xml_text = ET.tostring(document, encoding="unicode")
            if not character_pattern.search(xml_text) and "absPath" not in xml_text:
                continue
            for element in list(document.iter()):
                name = ET.QName(element).localname
                if name == "absPath":
                    element.getparent().remove(element)
                elif name == "cellStyle" and element.get("builtinId") in names:
                    element.set("name", names[element.get("builtinId")])
                elif name == "numFmt" and any(
                    x in element.get("formatCode", "") for x in ("mmm", "ddd")
                ):
                    date_format = element.get("formatCode")
                    if "[$-" not in date_format:
                        element.set("formatCode", "[$-040C]" + date_format)
                elif name == "theme" and character_pattern.search(element.get("name", "")):
                    element.set("name", "Office")
                elif name == "font" and element.get("script") in font_names:
                    font_name = element.get("typeface", "")
                    if character_pattern.search(font_name):
                        element.set(
                            "typeface",
                            font_names[element.get("script")]
                            + (" Light" if font_name.endswith(" Light") else ""),
                        )
            if path == "docProps/app.xml":
                labels = document.xpath("//*[local-name()='HeadingPairs']//*[local-name()='lpstr']")
                for element, name in zip(labels, ["Worksheets", "Named Ranges"]):
                    element.text = name
            changes[path] = ET.tostring(document, encoding="UTF-8", xml_declaration=True)
        if not changes:
            return
        temporary = file_path.with_name(file_path.stem + ".langue" + file_path.suffix)
        with ZipFile(temporary, "w", ZIP_DEFLATED) as output:
            for element in archive.infolist():
                output.writestr(element, changes.get(element.filename, archive.read(element.filename)))
    temporary.replace(file_path)


def sheet_paths(archive):
    workbook_xml = ET.fromstring(archive.read("xl/workbook.xml"))
    relations = {
        e.get("Id"): e.get("Target")
        for e in ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    }
    result = {}
    for sheet in workbook_xml.find("s:sheets", NS):
        path = relations[sheet.get("{" + NS["r"] + "}id")]
        result[sheet.get("name")] = (
            path.lstrip("/") if path.startswith("/") else "xl/" + path
        )
    return result


def preserve_inputs(source, destination):
    """Conserve les données de collecte et la présentation des graphiques."""
    # Un enregistrement Excel peut recalculer les entrées ou modifier les axes.
    # Reprendre ces éléments dans le fichier source conserve le modèle visuel.
    with ZipFile(source) as before, ZipFile(destination) as after:
        sheets_before, sheets_after = sheet_paths(before), sheet_paths(after)
        texts = []
        if "xl/sharedStrings.xml" in before.namelist():
            texts = [
                "".join(e.itertext())
                for e in ET.fromstring(before.read("xl/sharedStrings.xml"))
            ]
        changes = {}
        for name, path in sheets_before.items():
            if name in ("Dashboard", "Output", "Reco_histo"):
                continue
            data = deepcopy(ET.fromstring(before.read(path)).find("s:sheetData", NS))
            for cell in data.findall(".//s:c[@t='s']", NS):
                value = cell.find("s:v", NS)
                text = texts[int(value.text)]
                cell.remove(value)
                if cell.find("s:f", NS) is None:
                    cell.set("t", "inlineStr")
                    ET.SubElement(ET.SubElement(cell, S + "is"), S + "t").text = text
                else:
                    cell.set("t", "str")
                    ET.SubElement(cell, S + "v").text = text
            sheet = ET.fromstring(after.read(sheets_after[name]))
            sheet.replace(sheet.find("s:sheetData", NS), data)
            changes[sheets_after[name]] = ET.tostring(sheet, encoding="UTF-8", xml_declaration=True)
        # Excel peut modifier les axes pendant le chargement des données.
        for path in before.namelist():
            if not re.fullmatch(r"xl/charts/chart\d+\.xml", path):
                continue
            original = ET.fromstring(before.read(path))
            new = ET.fromstring(after.read(path))
            for cache in original.xpath("//*[local-name()='numCache' or local-name()='strCache']"):
                position = original.getroottree().getpath(cache)
                updated = new.xpath(position, namespaces={k: v for k, v in cache.nsmap.items() if k})
                if updated:
                    cache.getparent().replace(cache, deepcopy(updated[0]))
            changes[path] = ET.tostring(original, encoding="UTF-8", xml_declaration=True)
        for path in before.namelist():
            if re.fullmatch(r"xl/charts/(style|colors)\d+\.xml", path):
                changes[path] = before.read(path)
        workbook_xml = ET.fromstring(after.read("xl/workbook.xml"))
        calculation_settings = workbook_xml.find("s:calcPr", NS)
        if calculation_settings is not None:
            calculation_settings.set("calcMode", "auto")
            calculation_settings.set("fullCalcOnLoad", "0")
            calculation_settings.set("forceFullCalc", "0")
        changes["xl/workbook.xml"] = ET.tostring(workbook_xml, encoding="UTF-8", xml_declaration=True)
        for path in ("xl/_rels/workbook.xml.rels", "[Content_Types].xml"):
            document = ET.fromstring(after.read(path))
            for element in list(document):
                if "calcChain" in element.get("Type", "") or "calcChain" in element.get("PartName", ""):
                    document.remove(element)
            changes[path] = ET.tostring(document, encoding="UTF-8", xml_declaration=True)
        temporary = destination.with_suffix(".tmp.xlsm")
        with ZipFile(temporary, "w", ZIP_DEFLATED) as output:
            for element in after.infolist():
                if element.filename != "xl/calcChain.xml":
                    output.writestr(
                        element,
                        changes.get(element.filename, after.read(element.filename)),
                    )
    temporary.replace(destination)
